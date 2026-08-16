"""Create conservative, layout-preserving corrected paper drafts."""

from copy import deepcopy
import hashlib
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "ResearchPaper"
OUT = ROOT / "ResearchPaper_Revised"
OUT.mkdir(exist_ok=True)


def replace_runs(paragraph, new_text: str) -> None:
    """Replace text while preserving formatting of the paragraph's first run."""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def set_plain(paragraph, new_text: str) -> None:
    """Replace paragraph text and normalize inherited emphasis."""
    replace_runs(paragraph, new_text)
    for run in paragraph.runs:
        run.bold = False
        run.italic = False


def set_findings(paragraph, body: str) -> None:
    """Preserve the house style without making the entire finding emphatic."""
    for run in paragraph.runs:
        run.text = ""
    label = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    label.text = "Findings. "
    label.bold = True
    label.italic = True
    detail = paragraph.add_run(body)
    detail.bold = False
    detail.italic = False


def insert_before(paragraph, text: str, template) -> object:
    new_p = deepcopy(template._p)
    paragraph._p.addprevious(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    replace_runs(inserted, text)
    return inserted


def add_after(paragraph, text: str, template=None):
    source = template or paragraph
    new_p = deepcopy(source._p)
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    replace_runs(inserted, text)
    return inserted


def e1_metrics():
    values = {}
    for optimizer in ("adam_nv", "adam", "adamw", "adabelief"):
        seed_values = []
        for path in sorted((ROOT / "e1_convergence" / "logs").glob(f"{optimizer}_seed[0-9].csv")):
            frame = pd.read_csv(path)
            seed_values.append(float(frame["val_acc"].tail(5).mean()))
        series = pd.Series(seed_values, dtype=float)
        values[optimizer] = {
            "n": len(series),
            "mean": float(series.mean()),
            "sd": float(series.std(ddof=1)),
        }
    return values


def make_e1_figure(path: Path, metrics: dict) -> None:
    keys = ["adam_nv", "adam", "adamw", "adabelief"]
    labels = ["Adam-NV\n(n=4)", "Adam\n(n=2)", "AdamW\n(n=2)", "AdaBelief\n(n=2)"]
    means = [metrics[key]["mean"] for key in keys]
    errors = [metrics[key]["sd"] for key in keys]
    colors = ["#2b7fb8", "#888888", "#888888", "#888888"]
    fig, ax = plt.subplots(figsize=(7.0, 4.2))
    ax.bar(labels, means, yerr=errors, capsize=4, color=colors)
    ax.set_ylim(0.90, 0.945)
    ax.set_ylabel("Trailing-five-epoch validation accuracy")
    ax.set_title("E1 baseline convergence (provisional; unequal seed counts)")
    ax.grid(axis="y", alpha=0.2)
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)


def make_e11_figure(path: Path) -> None:
    """Plot controller states without inferring an unmeasured learning phase."""
    log_dir = ROOT / "e2_noisy_gradient" / "logs" / "Ablation"
    groups = {
        "Clean labels": sorted(log_dir.glob("adam_nv_sigma0.0_seed[0-9].csv")),
        "20% corrupted labels": sorted(log_dir.glob("adam_nv_sigma0.0_ln0.2_seed[0-9].csv")),
    }
    fig, axes = plt.subplots(1, 2, figsize=(8.8, 3.7), sharex=True)
    colors = {"Clean labels": "#2b7fb8", "20% corrupted labels": "#d06b38"}
    for label, paths in groups.items():
        frames = [pd.read_csv(item) for item in paths]
        epochs = frames[0]["epoch"]
        for axis, column, ylabel in zip(
            axes, ("beta1_mean", "sigma2_mean"), ("Mean beta_1", "Mean heuristic score")
        ):
            values = pd.concat([frame[column] for frame in frames], axis=1)
            mean = values.mean(axis=1)
            sd = values.std(axis=1, ddof=1)
            axis.plot(epochs, mean, label=label, color=colors[label], linewidth=2)
            axis.fill_between(epochs, mean - sd, mean + sd, color=colors[label], alpha=0.14)
            axis.set_ylabel(ylabel)
            axis.set_xlabel("Epoch")
            axis.grid(alpha=0.2)
            axis.axvline(26, color="#555555", linestyle=":", linewidth=1)
    axes[0].legend(frameon=False, fontsize=8)
    fig.suptitle("E11: Label noise — observed controller-state trajectories")
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)


def replace_embedded_image_by_hash(doc, original_path: Path, replacement_path: Path) -> bool:
    original_hash = hashlib.sha256(original_path.read_bytes()).digest()
    replacement = replacement_path.read_bytes()
    for relation in doc.part.rels.values():
        part = getattr(relation, "target_part", None)
        blob = getattr(part, "blob", None)
        if blob is not None and hashlib.sha256(blob).digest() == original_hash:
            part._blob = replacement
            return True
    return False


def revise_methodology() -> Path:
    source = PAPER / "IEEE_Access_Methodology.docx"
    target = OUT / "IEEE_Access_Methodology_Revised.docx"
    doc = Document(source)
    p = doc.paragraphs

    replace_runs(p[3],
        "Adam-NV (Noise-aware Variance-adaptive momentum) allows β₁ to vary "
        "per parameter element and per step as a function of a bounded heuristic "
        "noise score, denoted σ². The score combines the optimizer's prior "
        "first- and second-moment states and is warm-up rectified before being "
        "smoothed into rho. Because those moments can use different decay rates "
        "and are not placed on a common bias-corrected scale, sigma^2 is not "
        "interpreted as a calibrated variance or noise-to-signal estimator. The "
        "original controller lowers β₁ when the score rises and returns it "
        "toward its nominal value when the score falls. This design choice is "
        "tested explicitly through a reverse-direction smoothing ablation.")
    replace_runs(p[24],
        "Noise injection is performed after the backward pass and before the "
        "optimizer step. Matched seeds reproduce the same normalized Gaussian "
        "directions across optimizer runs. The final perturbation tensors are "
        "not byte-identical, because each direction is scaled by the gradient "
        "norm produced by that optimizer run.")
    replace_runs(p[29],
        "All runs log per-epoch training loss, training accuracy, validation "
        "loss, validation accuracy, learning rate, and wall-clock epoch duration "
        "to CSV. Adam-NV logs beta_1 and heuristic-score summaries. New runs "
        "also write JSON sidecars containing arguments, software versions, "
        "hardware details, the command, and RNG-stream seeds. Step-level response "
        "analysis uses diagnostic protocol v2, in which CosineAnnealingLR is "
        "advanced once per epoch, matching the main experiment.")
    replace_runs(p[37],
        "Hypothesis. Under the original controller, increasing injected noise is "
        "expected to increase the heuristic score and lower mean β₁. This is "
        "a mechanism hypothesis, not a claim that lower β₁ is the optimal "
        "smoothing response to noise.")
    replace_runs(p[42],
        "Design. Adam and Adam-NV were trained under the same four noise levels "
        "with matched model/data seeds and matched Gaussian perturbation "
        "directions; perturbation magnitudes remained run-specific because of "
        "gradient-norm scaling (Section 3.7).")
    replace_runs(p[54],
        "Design. Exploratory epoch durations were recorded during training. "
        "Because such values mix optimizer cost with data loading, forward and "
        "backward computation, and system variation, confirmatory overhead "
        "measurement uses repeated warm-up and timed runs, CUDA synchronization, "
        "separate optimizer-step and end-to-end timings, hardware/software "
        "metadata, and 95% confidence intervals.")
    replace_runs(p[59],
        "Metrics. Mean β₁ and heuristic-score trajectories, trailing-five-"
        "epoch validation accuracy, and degradation relative to clean training. "
        "Follow-up runs additionally measure training accuracy on uncorrupted "
        "examples and, on corrupted examples, accuracy against both the original "
        "and reassigned labels to test memorization directly.")
    replace_runs(p[66],
        "Mean β₁ and mean σ² serve as mechanism-level diagnostics. "
        "The σ² quantity is referred to as a heuristic noise score, not a calibrated "
        "variance estimate. Mean β₁ is the unweighted average across parameter "
        "tensors of each tensor's per-element mean and is not parameter-count "
        "weighted.")

    # Insert E13--E16 before Evaluation Metrics, preserving the E12 paragraph format.
    anchor = p[64]
    template_heading = p[60]
    template_body = p[61]
    additions = [
        "E13 — Non-stationary Gradient-Noise Schedules",
        "Goal. Test whether Adam-NV tracks changes within a run rather than only "
        "separating independent constant-noise conditions.",
        "Design. Adam and Adam-NV use matched seeds under phased, linear ramp-up, "
        "and intermittent schedules. The primary outcome is the trailing-five-"
        "epoch accuracy difference paired by seed; schedule-by-optimizer "
        "interactions are tested through direct contrasts of paired gaps.",
        "E14 — Coupled β₂ Sensitivity",
        "Goal. Screen whether shorter second-moment memory changes behavior under "
        "non-stationary noise, while acknowledging that β₂ jointly alters the "
        "noise heuristic and Adam's update denominator.",
        "Design. Adam-NV is evaluated at β₂ in {0.999, 0.99, 0.95} across "
        "constant, phased, ramp-up, and intermittent schedules with four matched seeds.",
        "E15 — Decoupled Noise-Memory Ablation",
        "Goal. Separate noise-tracking memory from the second-moment buffer used "
        "in the parameter update.",
        "Design. The update retains β₂=0.999 while an isolated noise-memory buffer "
        "uses decay values in {0.99, 0.95}. Accuracy contrasts and corrected "
        "protocol-v2 response times are paired by seed.",
        "E16 — CIFAR-100 Generality Evaluation",
        "Goal. Test whether findings extend beyond CIFAR-10.",
        "Design. Adam and Adam-NV are compared on CIFAR-100 under clean and "
        "intermittent-noise conditions using four matched seeds. This experiment "
        "is planned and must be completed before making cross-dataset claims.",
        "Controller-Direction Ablation",
        "Goal. Test the unresolved design assumption that higher heuristic noise "
        "should lower β₁. The original distrust rule is compared with a "
        "smoothing rule that raises β₁ from 0.9 toward 0.999 as the score rises.",
    ]
    current = anchor
    for index, text in enumerate(additions):
        template = template_heading if index in (0, 3, 6, 9, 12) else template_body
        current = insert_before(anchor, text, template)
        if index not in (0, 3, 6, 9, 12):
            set_plain(current, text)

    # Refresh paragraph collection after insertions and replace the final stats section.
    p = doc.paragraphs
    stats_heading = next(i for i, paragraph in enumerate(p) if paragraph.text.strip() == "3.11 Statistical Methodology")
    replace_runs(p[stats_heading + 1],
        "Each comparison uses matched random seeds where conditions share the "
        "same initialization, data-order, and perturbation-direction protocol. "
        "The primary accuracy outcome is the mean validation accuracy over the "
        "final five epochs. A comparison is publication-eligible only when all "
        "declared seed runs are complete.")
    replace_runs(p[stats_heading + 2],
        "Primary comparisons use paired mean differences with t-based 95% "
        "confidence intervals and paired t-tests. Exact two-sided sign-flip "
        "randomization p-values are reported as a robustness check for small "
        "samples. Holm correction controls family-wise error across comparisons "
        "within E13, E14, E15, and each other declared experiment family. Direct "
        "contrasts or interaction tests are required when comparing effects "
        "across schedules or designs; significance in one condition and not "
        "another is not itself evidence of a difference.")
    replace_runs(p[stats_heading + 3],
        "Null results are stated as no detectable advantage under the tested "
        "conditions and seed count. They are not treated as proof of equivalence "
        "or absence. Exploratory and incomplete analyses are labeled explicitly "
        "and excluded from confirmatory conclusions.")

    doc.save(target)
    return target


def revise_results() -> Path:
    source = PAPER / "IEEE_Access_Result.docx"
    target = OUT / "IEEE_Access_Result_Revised.docx"
    doc = Document(source)
    p = doc.paragraphs
    metrics = e1_metrics()
    figure_path = OUT / "e1_provisional.png"
    e11_figure_path = OUT / "e11_controller_trajectory.png"
    make_e1_figure(figure_path, metrics)
    make_e11_figure(e11_figure_path)

    set_plain(p[1],
        "This section reports the available findings in the order established in "
        "the methodology. E1 is provisional because baseline seed counts are "
        "unequal. E2, E7-A, and E7-C test controller response and adaptation; E9 "
        "reports exploratory timing; E11 and E12 test label and batch-size "
        "perturbations. E13–E15 are reanalyzed with paired statistics, while E16 "
        "and corrected protocol-v2 response-lag reruns remain pending. Findings "
        "are separated from interpretation, which is discussed in Section 5.")

    replace_runs(p[5],
        "Figure 4 shows provisional trailing-five-epoch validation accuracy. "
        "The repository contains four Adam-NV runs but only two complete runs "
        "for Adam, AdamW, and AdaBelief; inferential E1 comparisons are therefore "
        "withheld pending completion of baseline seeds 2 and 3.")
    replace_runs(p[8],
        "Figure 4. Provisional E1 comparison using available runs; unequal seed "
        "counts preclude confirmatory optimizer comparisons.")
    replace_runs(p[9],
        "Table 2. Provisional E1 baseline summary using available repository logs.")
    set_findings(p[12],
        "Adam-NV has four complete runs; Adam, AdamW, and AdaBelief each "
        "have two. The previously reported n=4 baseline table and significance "
        "claims are unsupported by the repository snapshot and are withdrawn. "
        "No E1 optimizer-superiority or equivalence claim should be made until "
        "seeds 2 and 3 are completed for every baseline and paired analysis is rerun.")

    # Replace first E1 image in place without changing drawing geometry.
    for shape in doc.inline_shapes:
        if shape._inline.docPr.get("name", "").startswith("Picture"):
            blip = shape._inline.graphic.graphicData.pic.blipFill.blip
            part = doc.part.related_parts[blip.embed]
            part._blob = figure_path.read_bytes()
            break

    table = doc.tables[0]
    display = {
        "Adam-NV (ours)": metrics["adam_nv"],
        "Adam": metrics["adam"],
        "AdamW": metrics["adamw"],
        "AdaBelief": metrics["adabelief"],
    }
    table.cell(0, 1).text = "Trailing-5 val. acc."
    table.cell(0, 2).text = "Available n"
    table.cell(0, 3).text = "Confirmatory comparison"
    for row in table.rows[1:]:
        label = row.cells[0].text.strip()
        value = display[label]
        row.cells[1].text = f"{value['mean']:.4f} ± {value['sd']:.4f}"
        row.cells[2].text = str(value["n"])
        row.cells[3].text = "Withheld - incomplete baselines"

    set_findings(p[80],
        "The displayed epoch-time ratios are exploratory and vary from "
        "0.988x to 1.562x across batch sizes. Because they were not obtained from "
        "repeated synchronized benchmark runs and lack uncertainty estimates, "
        "they do not support a stable-overhead claim. Confirmatory timing remains "
        "pending under the repeated warm-up and profiling protocol in Section 3.9.")
    replace_runs(p[85],
        "Figure 10 shows mean β₁ and heuristic-score trajectories for clean "
        "and 20%-label-corrupted training. The curves cross near epoch 26, but "
        "these controller trajectories alone cannot identify a memorization transition.")
    set_findings(p[92],
        "Mean β₁ and the heuristic score diverge between clean and "
        "label-corrupted training by epoch 60. The crossover near epoch 26 is a "
        "mechanism-state observation, not direct evidence of memorization. The "
        "claim remains provisional until training accuracy is measured separately "
        "on uncorrupted examples and corrupted examples against both clean and "
        "reassigned labels. No detectable robustness advantage was found in the "
        "reported interaction analysis.")
    set_plain(p[105],
        "Table 14 summarizes observed controller responses and performance outcomes "
        "for E2, E7-A, E7-C, E11, and E12. Section 4.9 then records the completion "
        "status of E13–E16 before Section 5 discusses the findings jointly.")
    set_plain(p[107],
        "Table 14. Summary of observed controller responses and performance outcomes.")
    doc.tables[8].cell(0, 1).text = "Logged controller response observed"
    replace_runs(p[112],
        "Across E2, E7-A, E7-C, E11, and E12, Adam-NV's logged controller state "
        "changed with the tested perturbations. These results validate activation "
        "of the implemented mechanism, not calibration of σ² as variance or "
        "optimality of lowering β₁ under noise. No tested condition established "
        "a detectable validation-accuracy advantage over Adam.")
    replace_runs(p[114],
        "The current evidence does not support Adam-NV as a general drop-in "
        "replacement for Adam for improved noisy-training accuracy. Earlier "
        "overhead values are exploratory; formal synchronized benchmarking is "
        "pending. Adam-NV is best viewed at present as a mechanism probe whose "
        "controller direction, generality, and runtime cost still require validation.")
    replace_runs(p[120],
        "The available E1 repository logs are incomplete for the baseline "
        "optimizers, so competitiveness claims are provisional. Across the "
        "completed noise-response experiments, no detectable accuracy advantage "
        "was found under the tested conditions.")
    replace_runs(p[132],
        "This study introduced Adam-NV, which varies β₁ per parameter element "
        "and step using a bounded heuristic noise score. The original rule lowers "
        "β₁ as the score rises. The score is not a calibrated variance or "
        "noise-to-signal estimator, and the optimal direction of the controller "
        "remains an empirical question.")
    replace_runs(p[133],
        "The reported CIFAR-10 experiments show that the controller state responds "
        "to several perturbation types. They do not establish optimizer superiority. "
        "E1 baseline comparisons remain provisional because only Adam-NV has four "
        "complete runs in the repository snapshot.")
    replace_runs(p[134],
        "No detectable improvement in final validation accuracy over Adam was found "
        "under the tested conditions and available seed counts. This wording does "
        "not imply equivalence or rule out smaller effects.")
    replace_runs(p[135],
        "Priority follow-up work is to complete missing baseline and E14 runs, "
        "rerun the corrected protocol-v2 response diagnostic across seeds, test "
        "the reverse controller direction, collect direct label-memorization "
        "metrics, benchmark overhead formally, and complete CIFAR-100 evaluation. "
        "E13-E15 statistical conclusions must use paired confidence intervals, "
        "direct contrasts, and Holm-adjusted p-values.")

    # Add an explicit status note before Discussion without fabricating unrun results.
    discussion = next(par for par in doc.paragraphs if par.text.strip() == "5. Discussion")
    insert_before(discussion, "4.9 E13–E16 — Extended Study Status", p[95])
    insert_before(discussion,
        "E13 non-stationary schedules, E14 coupled β₂ sensitivity, and E15 "
        "decoupled noise memory are present in the repository and have been "
        "reanalyzed with paired seed matching. In the current exploratory table, "
        "no E13 Adam-NV-versus-Adam contrast and no E15 decoupled-versus-coupled "
        "contrast survives Holm correction. E14 suggests a ramp-up degradation at "
        "β₂=0.95 in the adjusted paired t-test, but the exact four-seed sign-"
        "flip p-value is 0.125. One intermittent E14 run ends at epoch 32, so the "
        "analysis correctly refuses publication-mode output. E16 CIFAR-100 and all "
        "protocol-v2 lag reruns remain pending. These extended results must replace, "
        "not be combined with, the invalid legacy 1,317-to-47-step lag claim.",
        p[2])

    source_e11 = ROOT / "e2_noisy_gradient" / "figures" / "e11_crossover_figure.png"
    if not replace_embedded_image_by_hash(doc, source_e11, e11_figure_path):
        raise RuntimeError("Could not locate the embedded E11 figure")

    doc.save(target)
    return target


if __name__ == "__main__":
    print(revise_methodology())
    print(revise_results())
